#!/usr/bin/env python
"""把论文 Markdown 稿导出为 Word（.docx）。

为什么不是简单地把 Markdown 贴进 Word
--------------------------------------
学术稿件里公式是第一等公民。纯文本转换会把公式变成不可读的 LaTeX 源码，
或者变成无法再编辑的图片。本脚本走的是:

    LaTeX --latex2mathml--> MathML --MML2OMML.XSL--> OMML

OMML 是 Word 原生的公式格式，生成的是**可双击编辑的真实公式对象**，
与在 Word 里用公式编辑器敲出来的一样。XSLT 样式表由 Microsoft Office
自带（``MML2OMML.XSL``），脚本会自动定位。

格式
----
* 正文 Times New Roman 11 pt，行距 1.4，两端对齐
* A4 纸，页边距 2.5 cm
* 公式居中、独立成段
* 表格用三线表样式（学术论文惯例）
* 待补数字标记 ``⟨PENDING⟩`` 加黄色高亮，避免误当成已完成内容

用法
----
::

    python scripts/export_docx.py                    # 默认导出 manuscript_draft.md
    python scripts/export_docx.py --input paper/xxx.md --output paper/xxx.docx
"""

from __future__ import annotations

import argparse
import logging
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from evtol_siting.config import PROJECT_ROOT  # noqa: E402
from evtol_siting.logging_setup import setup_logging  # noqa: E402

log = logging.getLogger("export_docx")

# Office 自带的 MathML -> OMML 样式表候选路径
XSLT_CANDIDATES = [
    r"C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL",
    r"C:\Program Files (x86)\Microsoft Office\root\Office16\MML2OMML.XSL",
    r"C:\Program Files\Microsoft Office\Office16\MML2OMML.XSL",
    "/Applications/Microsoft Word.app/Contents/Resources/MML2OMML.XSL",
    "/usr/share/office/MML2OMML.XSL",
]


def find_xslt() -> str | None:
    """定位 Office 自带的 MML2OMML.XSL。"""
    import os

    for p in XSLT_CANDIDATES:
        if os.path.exists(p):
            return p
    # 兜底：在 Office 安装目录下搜一层
    for root in (r"C:\Program Files\Microsoft Office", r"C:\Program Files (x86)\Microsoft Office"):
        if not os.path.isdir(root):
            continue
        for dirpath, _dirnames, filenames in os.walk(root):
            if "MML2OMML.XSL" in filenames:
                return os.path.join(dirpath, "MML2OMML.XSL")
    return None


class MathConverter:
    """LaTeX -> OMML 转换器（带缓存，同一公式只转一次）。"""

    def __init__(self):
        from lxml import etree

        self._cache: dict[str, bytes] = {}
        self._enabled = False
        xslt_path = find_xslt()
        if xslt_path:
            try:
                self._transform = etree.XSLT(etree.parse(xslt_path))
                self._etree = etree
                self._enabled = True
                log.info("公式转换已启用，样式表: %s", xslt_path)
            except Exception as exc:
                log.warning("加载 XSLT 失败: %s", str(exc)[:150])
        if not self._enabled:
            log.warning(
                "未找到 MML2OMML.XSL —— 公式将退化为纯文本。"
                "安装 Microsoft Office 后重跑即可获得可编辑的 Word 公式对象。"
            )

    @property
    def enabled(self) -> bool:
        return self._enabled

    def to_omml(self, latex: str) -> bytes | None:
        """LaTeX 片段 -> OMML 字节串；失败返回 None（调用方退化为文本）。"""
        if not self._enabled:
            return None
        latex = latex.strip()
        if latex in self._cache:
            return self._cache[latex]
        try:
            import latex2mathml.converter as conv

            mathml = conv.convert(latex)
            omml = self._etree.tostring(
                self._transform(self._etree.fromstring(mathml.encode()))
            )
            if b"oMath" not in omml:
                return None
            self._cache[latex] = omml
            return omml
        except Exception as exc:
            log.debug("公式转换失败 %r: %s", latex[:60], str(exc)[:100])
            return None


# ---------------------------------------------------------------------------
# 行内标记解析
# ---------------------------------------------------------------------------

# 把一段文本拆成 (类型, 内容) 序列：text / math / bold / italic / code
_INLINE_RE = re.compile(
    r"(\$[^$]+\$)"          # 行内公式
    r"|(\*\*[^*]+\*\*)"     # 粗体
    r"|(\*[^*]+\*)"         # 斜体
    r"|(`[^`]+`)"           # 等宽
)


def split_inline(text: str) -> list[tuple[str, str]]:
    """把一行文本拆成带类型标记的片段。"""
    out: list[tuple[str, str]] = []
    pos = 0
    for m in _INLINE_RE.finditer(text):
        if m.start() > pos:
            out.append(("text", text[pos:m.start()]))
        tok = m.group(0)
        if tok.startswith("$"):
            out.append(("math", tok[1:-1]))
        elif tok.startswith("**"):
            out.append(("bold", tok[2:-2]))
        elif tok.startswith("`"):
            out.append(("code", tok[1:-1]))
        else:
            out.append(("italic", tok[1:-1]))
        pos = m.end()
    if pos < len(text):
        out.append(("text", text[pos:]))
    return out


# ---------------------------------------------------------------------------
# 文档构建
# ---------------------------------------------------------------------------

def add_runs(par, text: str, math: MathConverter, base_bold=False, base_italic=False):
    """向段落写入带格式的文本片段；公式片段插入 OMML 对象。"""
    from docx.shared import Pt

    for kind, content in split_inline(text):
        if kind == "math":
            omml = math.to_omml(content)
            if omml is not None:
                par._p.append(math._etree.fromstring(omml))
                continue
            # 退化：用 Unicode 化的文本表示
            r = par.add_run(_latex_fallback(content))
            r.font.name = "Cambria Math"
            r.font.size = Pt(11)
            r.italic = True
            continue

        r = par.add_run(content)
        r.bold = base_bold or (kind == "bold")
        r.italic = base_italic or (kind == "italic")
        if kind == "code":
            r.font.name = "Consolas"
            r.font.size = Pt(10)


_GREEK = {
    r"\mu": "μ", r"\sigma": "σ", r"\lambda": "λ", r"\beta": "β", r"\rho": "ρ",
    r"\tau": "τ", r"\alpha": "α", r"\gamma": "γ", r"\delta": "δ", r"\epsilon": "ε",
    r"\sum": "Σ", r"\in": "∈", r"\ge": "≥", r"\le": "≤", r"\times": "×",
    r"\cdot": "·", r"\infty": "∞", r"\forall": "∀", r"\ni": "∋",
    r"\max": "max", r"\min": "min", r"\exp": "exp", r"\log": "log",
    r"\left": "", r"\right": "", r"\quad": "  ", r"\,": " ",
}


def _latex_fallback(latex: str) -> str:
    """在没有 XSLT 时，把 LaTeX 粗略转成可读的 Unicode 文本。"""
    s = latex
    for k, v in _GREEK.items():
        s = s.replace(k, v)
    s = s.replace(r"\frac{", "(").replace("}{", ")/(").replace("}", ")")
    s = s.replace("_", "").replace("^", "^").replace("\\", "")
    return s


def add_table(doc, rows: list[list[str]], math: MathConverter):
    """写一个三线表（学术论文惯例：仅顶线、表头下线、底线）。"""
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt

    if not rows:
        return
    n_cols = max(len(r) for r in rows)
    rows = [r + [""] * (n_cols - len(r)) for r in rows]

    tbl = doc.add_table(rows=len(rows), cols=n_cols)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = True

    for i, row in enumerate(rows):
        for j, cell_text in enumerate(row):
            cell = tbl.cell(i, j)
            par = cell.paragraphs[0]
            par.paragraph_format.space_before = Pt(2)
            par.paragraph_format.space_after = Pt(2)
            add_runs(par, cell_text.strip(), math, base_bold=(i == 0))
            for r in par.runs:
                r.font.size = Pt(9.5)

    # 三线表边框：去掉全部框线，再加顶线、表头下线、底线
    def _set_border(cell, edge, sz, val="single"):
        tcPr = cell._tc.get_or_add_tcPr()
        borders = tcPr.find(qn("w:tcBorders"))
        if borders is None:
            borders = OxmlElement("w:tcBorders")
            tcPr.append(borders)
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), val)
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:color"), "000000")

    for j in range(n_cols):
        _set_border(tbl.cell(0, j), "top", 12)
        _set_border(tbl.cell(0, j), "bottom", 6)
        _set_border(tbl.cell(len(rows) - 1, j), "bottom", 12)


def build_docx(md_path: Path, out_path: Path, style: str = "default") -> Path:
    """把 Markdown 稿件转成 Word 文档。"""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
    from docx.shared import Cm, Pt, RGBColor

    math = MathConverter()
    md = md_path.read_text(encoding="utf-8")

    doc = Document()

    # -- 页面与正文样式 ----------------------------------------------------
    #
    # ``--style mdpi`` 套用 MDPI（Aerospace）排版。规格不是猜的，取自
    # 本刊一篇已发表论文的实测值（论文/GIS/aerospace-12-00709 (1).pdf）：
    #   A4；正文 Palatino 10 pt；标题 18 pt；一级标题 12 pt；图注 8 pt；
    #   左右边距实测 1.25 / 1.21 cm（与 MDPI Layout Style Guide 的 1.27 cm 吻合），
    #   上下边距按规范取 2.5 / 1.9 cm。
    # MDPI 用 URW Palladio L（Palatino 的自由克隆）；在 Word 中以
    # "Palatino Linotype" 最接近，缺失时回退到 Book Antiqua / 宋体。
    is_mdpi = str(style).lower() == "mdpi"

    sec = doc.sections[0]
    # python-docx 的空白模板默认是 US Letter，必须显式设成 A4——
    # 否则 MDPI 要求 A4 而导出的是 Letter，投稿系统会退回来。
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    if is_mdpi:
        sec.top_margin = Cm(2.5)
        sec.bottom_margin = Cm(1.9)
        sec.left_margin = sec.right_margin = Cm(1.27)
    else:
        sec.top_margin = sec.bottom_margin = Cm(2.5)
        sec.left_margin = sec.right_margin = Cm(2.5)

    # 判断稿件语言：中文稿需要把 eastAsia 字体设为中文字体，否则 Word 会用
    # 默认的宋体渲染，标题也会失去层级感。西文部分的字体（Times New Roman）
    # 仍通过 ascii/hAnsi 指定，二者在 Word 中是分开的槽位。
    is_zh = len(re.findall(r"[一-鿿]", md[:20000])) > 500
    if is_mdpi:
        latin = "Palatino Linotype"
        cjk_body = "宋体"
        cjk_head = "黑体"
    else:
        latin = "Times New Roman"
        cjk_body = "宋体"
        cjk_head = "黑体" if is_zh else "宋体"
    log.info("稿件语言: %s", "中文" if is_zh else "英文")

    def _set_fonts(style, latin_name, cjk_name):
        from docx.oxml.ns import qn

        style.font.name = latin_name
        rpr = style.element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            from docx.oxml import OxmlElement

            rfonts = OxmlElement("w:rFonts")
            rpr.append(rfonts)
        # 显式字体属性与主题字体属性（*Theme）在 OOXML 中互斥。
        # Word 内置的 Heading 样式默认带 asciiTheme/eastAsiaTheme，
        # 若不删除，某些 Word 版本会优先采用主题字体、忽略这里的显式设置。
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            if rfonts.get(qn(attr)) is not None:
                del rfonts.attrib[qn(attr)]
        rfonts.set(qn("w:ascii"), latin_name)
        rfonts.set(qn("w:hAnsi"), latin_name)
        rfonts.set(qn("w:eastAsia"), cjk_name)

    normal = doc.styles["Normal"]
    normal.font.size = Pt(10 if is_mdpi else 11)
    _set_fonts(normal, latin, cjk_body)
    pf = normal.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    if is_mdpi:
        # MDPI 排版紧凑：单倍行距，段后 4 pt，首行不缩进
        pf.line_spacing = 1.08
        pf.space_after = Pt(4)
    else:
        # 中文行距习惯略大于西文
        pf.line_spacing = 1.5 if is_zh else 1.4
        pf.space_after = Pt(6)

    # MDPI 标题字号实测：一级 12 pt；二三级依次递减，均加粗
    head_sizes = ((1, 12), (2, 11), (3, 10), (4, 10)) if is_mdpi \
        else ((1, 15), (2, 13), (3, 12), (4, 11))
    for lvl, size in head_sizes:
        st = doc.styles[f"Heading {lvl}"]
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor(0, 0, 0)
        _set_fonts(st, latin, cjk_head)

    lines = md.split("\n")
    i = 0
    in_table = False
    table_rows: list[list[str]] = []
    n_equations = 0
    n_converted = 0
    n_images = 0
    n_missing_images = 0
    pending_caption = False

    def flush_table():
        nonlocal in_table, table_rows
        if table_rows:
            add_table(doc, table_rows, math)
            doc.add_paragraph()
        in_table, table_rows = False, []

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # -- 图件 ![alt](path) ---------------------------------------------
        m_img = re.match(r"^!\[.*?\]\((.+?)\)\s*$", stripped)
        if m_img:
            from docx.shared import Cm

            rel = m_img.group(1).strip()
            # 路径相对论文文件夹解析（图片与稿件同目录树）
            cand = (md_path.parent / rel)
            if not cand.exists():
                cand = md_path.parent / Path(rel).name
            if cand.exists():
                par = doc.add_paragraph()
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                par.paragraph_format.space_before = Pt(8)
                par.paragraph_format.space_after = Pt(2)
                try:
                    par.add_run().add_picture(str(cand), width=Cm(15.5))
                    n_images += 1
                    pending_caption = True
                except Exception as exc:
                    log.warning("插入图件失败 %s: %s", cand.name, str(exc)[:120])
                    n_missing_images += 1
            else:
                log.warning("图件不存在，已跳过: %s", rel)
                n_missing_images += 1
                par = doc.add_paragraph()
                r = par.add_run(f"[缺失图件: {rel}]")
                r.italic = True
            i += 1
            continue

        # -- 表格 ----------------------------------------------------------
        if stripped.startswith("|") and stripped.endswith("|"):
            cells = [c for c in stripped.strip("|").split("|")]
            # 跳过分隔行 |---|---|
            if not re.fullmatch(r"[\s:\-|]+", stripped):
                table_rows.append(cells)
            in_table = True
            i += 1
            continue
        elif in_table:
            flush_table()

        # -- 独立公式 $$ ... $$ ---------------------------------------------
        if stripped.startswith("$$"):
            body = stripped[2:]
            if body.endswith("$$"):
                body = body[:-2]
            else:
                i += 1
                while i < len(lines) and not lines[i].strip().endswith("$$"):
                    body += "\n" + lines[i]
                    i += 1
                if i < len(lines):
                    body += "\n" + lines[i].strip()[:-2]
            par = doc.add_paragraph()
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.space_before = Pt(6)
            par.paragraph_format.space_after = Pt(6)
            n_equations += 1
            omml = math.to_omml(body.strip())
            if omml is not None:
                n_converted += 1
                par._p.append(math._etree.fromstring(omml))
            else:
                r = par.add_run(_latex_fallback(body.strip()))
                r.font.name = "Cambria Math"
                r.italic = True
            i += 1
            continue

        # -- 水平线 --------------------------------------------------------
        if re.fullmatch(r"-{3,}", stripped):
            i += 1
            continue

        # -- 标题 ----------------------------------------------------------
        m = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if m:
            lvl = len(m.group(1))
            h = doc.add_heading(level=lvl)
            h.paragraph_format.space_before = Pt(14 if lvl <= 2 else 10)
            h.paragraph_format.space_after = Pt(6)
            add_runs(h, m.group(2), math)
            for r in h.runs:
                r.font.color.rgb = RGBColor(0, 0, 0)
            i += 1
            continue

        # -- 引用块 --------------------------------------------------------
        if stripped.startswith(">"):
            par = doc.add_paragraph()
            par.paragraph_format.left_indent = Cm(0.8)
            par.paragraph_format.right_indent = Cm(0.8)
            add_runs(par, stripped.lstrip("> ").strip(), math, base_italic=True)
            for r in par.runs:
                r.font.size = Pt(10)
            i += 1
            continue

        # -- 列表 ----------------------------------------------------------
        m = re.match(r"^(\s*)[-*]\s+(.*)$", line)
        if m:
            par = doc.add_paragraph(style="List Bullet")
            par.paragraph_format.space_after = Pt(3)
            add_runs(par, m.group(2), math)
            i += 1
            continue
        m = re.match(r"^(\s*)\d+\.\s+(.*)$", line)
        if m:
            par = doc.add_paragraph(style="List Number")
            par.paragraph_format.space_after = Pt(3)
            add_runs(par, m.group(2), math)
            i += 1
            continue

        # -- 空行 ----------------------------------------------------------
        if not stripped:
            i += 1
            continue

        # -- 图注（紧随图片之后的那一段）--------------------------------------
        # ⚠ 必须同时匹配英文 `**Figure` 与中文 `**图`。早期只匹配前者，
        #   于是**中文稿的全部图注都没有按图注格式排版**（无 9pt 字号、
        #   无左右缩进）——docx XML 里可直接验证，而 Markdown 源码看不出差别。
        if pending_caption and re.match(r"^\*\*(Figure|图)", stripped):
            par = doc.add_paragraph()
            par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            par.paragraph_format.left_indent = Cm(0.6)
            par.paragraph_format.right_indent = Cm(0.6)
            par.paragraph_format.space_after = Pt(10)
            add_runs(par, stripped, math)
            for r in par.runs:
                r.font.size = Pt(9)
            pending_caption = False
            i += 1
            continue
        pending_caption = False

        # -- 正文 ----------------------------------------------------------
        par = doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_runs(par, stripped, math)

        # 待办标记加高亮，避免误读为已完成
        if "PENDING" in stripped or "DRAFT" in stripped:
            from docx.enum.text import WD_COLOR_INDEX

            for r in par.runs:
                if "PENDING" in r.text or "DRAFT" in r.text or "⟨" in r.text:
                    r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        i += 1

    flush_table()

    # -- 页脚页码 ----------------------------------------------------------
    try:
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn

        footer = doc.sections[0].footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = fp.add_run()
        for instr in ("begin", "PAGE", "end"):
            el = OxmlElement("w:fldChar") if instr != "PAGE" else OxmlElement("w:instrText")
            if instr == "PAGE":
                el.set(qn("xml:space"), "preserve")
                el.text = " PAGE "
            else:
                el.set(qn("w:fldCharType"), instr)
            run._r.append(el)
    except Exception:
        pass

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)

    log.info(
        "已导出 %s\n    公式 %d 个（其中 %d 个转为 Word 原生公式对象）\n    表格已转为三线表",
        out_path, n_equations, n_converted,
    )
    if math.enabled and n_converted < n_equations:
        log.warning("%d 个公式转换失败，已退化为文本", n_equations - n_converted)
    return out_path


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="论文 Markdown -> Word")
    ap.add_argument("--input", default=None,
                    help="输入 Markdown（默认取 ../论文稿件/manuscript_draft.md）")
    ap.add_argument("--output", default=None, help="输出 .docx")
    ap.add_argument("--style", choices=["default", "mdpi"], default="default",
                    help="排版预设。mdpi = 套用 MDPI（Aerospace）规格："
                         "A4、Palatino 10 pt、标题 12 pt、页边距 2.5/1.9/1.27 cm。"
                         "规格取自本刊已发表论文的实测值，见脚本内注释。")
    ap.add_argument("-v", "--verbose", action="store_true")
    args = ap.parse_args(argv)

    setup_logging(args.verbose, PROJECT_ROOT / "outputs" / "logs" / "export_docx.log")

    # ⚠ 默认输入曾是 `paper/manuscript_draft.md`——那是**早已过期的副本**
    #   （2026-09-15 的 51 KB 版，而正稿在 `论文稿件/` 已 180 KB+）。
    #   该默认值造成过一次实际后果：一次不带 --input 的导出把旧稿渲染成了
    #   `paper/manuscript_draft.docx`，看上去与正稿无异。
    #   正稿的唯一位置是 `论文稿件/`，故默认值改指那里，并在取不到时报错而非回退。
    _default = PROJECT_ROOT.parent / "论文稿件" / "manuscript_draft.md"
    src = Path(args.input) if args.input else _default
    if not src.exists():
        raise SystemExit(
            f"找不到默认输入 {src}。\n"
            "论文正文的唯一权威副本在 `论文稿件/`；请用 --input 显式指定，"
            "不要回退到任何 `paper/` 下的副本。")
    dst = Path(args.output) if args.output else src.with_suffix(".docx")
    if not src.exists():
        log.error("输入文件不存在: %s", src)
        return 1

    build_docx(src, dst, style=args.style)
    print(f"\n已生成: {dst}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
